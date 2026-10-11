from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).with_name("Assignment_02_Wan_Jiahao.docx")
BLACK = RGBColor(0, 0, 0)
BODY_FONT = "Times New Roman"
EAST_ASIA_BODY = "宋体"
HEADING_FONT = "SimHei"


def set_run(run, *, heading=False, bold=False, italic=False, size=None):
    run.font.name = HEADING_FONT if heading else BODY_FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), HEADING_FONT if heading else EAST_ASIA_BODY)
    run.font.size = Pt(size or (14 if heading else 12))
    run.font.color.rgb = BLACK
    run.bold = bold
    run.italic = italic


def set_paragraph(paragraph, *, after=6, before=0, line=1.15, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_after = Pt(after)
    fmt.space_before = Pt(before)
    fmt.line_spacing = line
    fmt.keep_with_next = keep


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    set_paragraph(paragraph, before=8 if level == 1 else 5, after=4, keep=True)
    run = paragraph.add_run(text)
    set_run(run, heading=True, bold=True)
    return paragraph


def add_body(doc, text, *, bold_lead=None, keep=False):
    paragraph = doc.add_paragraph()
    set_paragraph(paragraph, keep=keep)
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run(lead, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_run(rest)
    else:
        run = paragraph.add_run(text)
        set_run(run)
    return paragraph


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def three_line_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = 1
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "12")
        item.set(qn("w:color"), "000000")
        borders.append(item)
    for edge in ("left", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "nil")
        borders.append(item)
    tbl_pr.append(borders)

    header_row = table.rows[0]
    tr_pr = header_row._tr.get_or_add_trPr()
    header_border = OxmlElement("w:trBorders")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:color"), "000000")
    header_border.append(bottom)
    tr_pr.append(header_border)

    for index, text in enumerate(headers):
        cell = header_row.cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        set_paragraph(p, after=0, line=1.0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run(run, bold=True)

    for values in rows:
        cells = table.add_row().cells
        for index, text in enumerate(values):
            cell = cells[index]
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            set_paragraph(p, after=0, line=1.0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index < 2 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            set_run(run)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run(run, size=10)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.85)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

normal = doc.styles["Normal"]
normal.font.name = BODY_FONT
normal._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA_BODY)
normal.font.size = Pt(12)
normal.font.color.rgb = BLACK
for style_name in ("Title", "Heading 1", "Heading 2"):
    style = doc.styles[style_name]
    style.font.name = HEADING_FONT
    style._element.rPr.rFonts.set(qn("w:eastAsia"), HEADING_FONT)
    style.font.size = Pt(14)
    style.font.color.rgb = BLACK
    style.font.bold = True
    style_ppr = style.element.get_or_add_pPr()
    existing_border = style_ppr.find(qn("w:pBdr"))
    if existing_border is not None:
        style_ppr.remove(existing_border)

# Cover sheet
for _ in range(3):
    doc.add_paragraph()
cover = doc.add_paragraph()
cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_paragraph(cover, after=18)
run = cover.add_run("GEC Academy\nOnline Research Seminar")
set_run(run, heading=True, bold=True)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_paragraph(title, before=16, after=24)
run = title.add_run("Assignment Lab 2")
set_run(run, heading=True, bold=True)

cover_rows = [
    ("Name", "Wan, Jiahao"),
    ("Date", "October 11, 2026"),
    ("Email Address", "2621751515@qq.com"),
    ("Course", "Introduction to Ethereum and Cloud Blockchain Solutions Engineering"),
    ("Repository", "https://github.com/onemoretime-hh/ethereum-cloud-blockchain-labs"),
]
table = doc.add_table(rows=0, cols=2)
table.alignment = 1
table.autofit = False
for label, value in cover_rows:
    cells = table.add_row().cells
    cells[0].width = Inches(1.55)
    cells[1].width = Inches(4.85)
    for cell in cells:
        set_cell_margins(cell, top=100, bottom=100)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            item = OxmlElement(f"w:{edge}")
            item.set(qn("w:val"), "nil")
            borders.append(item)
        tc_pr.append(borders)
    p0 = cells[0].paragraphs[0]
    p1 = cells[1].paragraphs[0]
    set_paragraph(p0, after=0)
    set_paragraph(p1, after=0)
    r0 = p0.add_run(label)
    r1 = p1.add_run(value)
    set_run(r0, bold=True)
    set_run(r1)

doc.add_paragraph()
statement = add_body(doc, "Statement of independent effort: I affirm that this submission represents my own independent effort and that all sources used are identified in the references.")
statement.paragraph_format.space_before = Pt(20)
add_body(doc, "Electronic signature: Jiahao Wan")
add_body(doc, "Total in points (100 points total): ____________________")
add_body(doc, "Professor's comments: _________________________________________________")

doc.add_page_break()

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_paragraph(title, after=12)
run = title.add_run("Assignment Answer Sheet")
set_run(run, heading=True, bold=True)
add_body(doc, "This report answers the Ethereum exercises and documents a compact modern implementation of the Casino decentralized application. The implementation is intended for Sepolia classroom testing and demonstrates the complete path from a Solidity contract to a browser interface.")

add_heading(doc, "Part I Exercises", 1)
add_heading(doc, "1 Smart contract use cases", 2)
add_body(doc, "1(a) The lecture identifies four representative uses. Token systems use contracts to issue, transfer, and account for digital units under a shared rule set. Identity and reputation systems associate names, credentials, or reputation records with blockchain-controlled accounts. Decentralized autonomous organizations use contracts to hold assets and enforce collectively approved rules. Election and voting systems record eligible choices and apply transparent counting rules so that the result can be independently checked.")
add_body(doc, "1(b) A decentralized exchange is another use case: liquidity and swap contracts execute asset exchanges without a central matching operator. Decentralized lending is a second use case: contracts accept supplied assets, calculate debt, hold collateral, and enforce liquidation rules. A third use case is an NFT marketplace, where contracts represent unique assets and coordinate offers, sales, royalties, and ownership transfers.")
add_body(doc, "1(c) The following Ethereum mainnet contracts are concrete examples. Addresses were checked against the respective protocol documentation.")
three_line_table(
    doc,
    ["Use case", "Contract", "Ethereum mainnet address", "Reason for classification"],
    [
        ("Token system", "USDC", "0xA0b86991c6218b36c1d19D4a2e9Eb0cE3606eB48", "The contract controls balances and transfers of the USDC token."),
        ("Identity", "ENS Registry", "0x00000000000C2E074eC69A0dFb2997BA6C7d2e1e", "The registry is the starting point for ownership and resolution of human-readable ENS names."),
        ("DAO and voting", "Aave Governance", "0x9AEE0B04504CeF83A65AC3f0e838D0593BCb2BC7", "The governance contract coordinates protocol proposals and on-chain voting decisions."),
        ("Exchange", "Uniswap V2 Router02", "0x7a250d5630B4cF539739dF2C5dAcb4c659F2488D", "The router exposes token-swap and liquidity-routing operations for Uniswap V2."),
        ("Lending", "Aave V3 Pool", "0x87870Bca3F3fD6335C3F4ce8392D69350B4fA4E2", "The pool is the principal user-facing contract for supplying, borrowing, repaying, and withdrawing assets."),
    ],
    [1.0, 1.05, 2.05, 2.7],
)

add_heading(doc, "2 Gas and EVM instruction costs", 2)
add_body(doc, "2(a) Gas is Ethereum's unit for measuring the computational and storage work performed by the EVM. Each operation consumes a defined amount of gas. A transaction fee is based on the gas actually used and the applicable fee per gas. This metering limits resource consumption and compensates the network for processing state-changing transactions.")
add_body(doc, "2(b) The sender specifies a gas limit, but the account is charged only for gas actually consumed, subject to the transaction fee rules. Gas that remains when execution ends successfully is not spent and its corresponding value remains with, or is returned to, the sender. The gas limit is therefore a maximum allowance rather than an automatic charge for every reserved unit.")
add_body(doc, "2(c) Under the current schedule, and comparing the stated base cases without memory expansion, transferred value, account warming penalties, forwarded gas, or code-deposit cost, the ascending order is shown below. PUSH1 and PUSH4 have the same base cost. Dynamic components can change the total cost of SHA3, CALL, SELFDESTRUCT, and CREATE in a real execution.")
three_line_table(
    doc,
    ["Rank", "Instruction", "Base gas used for comparison", "Qualification"],
    [
        ("1", "STOP", "0", "Ends execution."),
        ("2", "PUSH1 0xff", "3", "Push instructions share the very-low base tier."),
        ("2", "PUSH4 0x00000000", "3", "Equal to PUSH1 in base gas."),
        ("3", "JUMP", "8", "Destination must be a valid JUMPDEST."),
        ("4", "SHA3 / KECCAK256", "30 plus 6 per word", "Memory expansion may add cost."),
        ("5", "CALL", "100 warm base", "Cold access, value transfer, memory, and forwarded gas can add cost."),
        ("6", "SELFDESTRUCT", "5,000 base", "The question assumes that no new account is created; access cost can still apply."),
        ("7", "CREATE", "32,000 base", "Init-code execution, memory expansion, and code deposit add cost."),
    ],
    [0.55, 1.45, 1.55, 3.25],
)

add_heading(doc, "3 Prevention of double spending", 2)
add_body(doc, "Ethereum prevents double spending through authenticated state transitions and per-account nonces. Nodes verify each transaction's signature, balance, chain rules, and expected nonce. Only one transaction with a particular nonce can execute for an account. Inclusion updates the nonce and balances in the agreed world state, so a competing transaction with the same nonce becomes invalid and a later transaction cannot overspend the remaining balance. Proof-of-stake consensus gives nodes a common canonical ordering of these changes.")

add_heading(doc, "4 Transactions and internal messages", 2)
add_body(doc, "A contract-to-contract message is not an independent, signed transaction. It is an execution step produced while the EVM processes an outer transaction. Publishing every internal call as another blockchain transaction would duplicate information because the outer transaction, the prior world state, and the contract bytecode already determine every call, return value, log, and state change. Each validating node executes the same ordered transactions with the same deterministic EVM rules. Consequently, nodes reconstruct the same internal calls and arrive at the same state root without storing those calls as separate consensus objects. Explorers may index call traces and label them internal transactions, but those traces are derived data rather than standalone transactions in blocks.")

add_heading(doc, "5 Transactions that remain pending or fail", 2)
add_body(doc, "5(a) Four common reasons a submitted transaction may not be mined are an effective fee that is too low for validators to select it, a nonce gap caused by an earlier pending transaction, replacement by another transaction from the same account with the same nonce and a higher fee, and insufficient ETH to cover the transferred value plus the maximum fee. A node may also drop a transaction when its pool is full or when propagation is poor, but four reasons are sufficient here.")
add_body(doc, "5(b) A transaction that has been included can fail because execution runs out of gas. It can also fail because contract logic reverts, for example when a require condition is false; an invalid opcode or other exceptional halt has a similar failed status.")
add_body(doc, "5(c) A failed transaction still costs gas because validators executed and verified it, consumed computation and bandwidth, and included it in a block. Reverting rolls back the contract's state changes, but it does not undo the resources already used. Charging for those resources also prevents attackers from forcing unlimited failing computation at no cost.")

add_heading(doc, "Part II Modern Casino DApp", 1)
add_heading(doc, "Implementation overview", 2)
add_body(doc, "The project implements the tutorial's central exercise in a deliberately small form. Casino.sol accepts one bet per address per round on a number from 1 through 10. The constructor sets the minimum stake and the number of bets needed to close a round; the deployment script uses 0.001 ETH and 100 bets. After a round closes, settleRound selects a teaching-only winning number, credits each winner with an equal share of the prize pool, and opens the next round. If no bet selected the winning number, the pot rolls into the next round. Winners withdraw through claimWinnings.")
add_body(doc, "The browser application is a single React page. It requests authorized accounts with the current MetaMask provider flow, checks for Sepolia chain ID 11155111, creates an ethers.js BrowserProvider, and connects the deployed contract through its ABI. The interface displays the current round, number of bets, pot, and round status. It provides controls to place a bet, settle a full round, and claim winnings. Buttons are disabled while a transaction is pending, and status text reports connection or transaction errors.")
add_body(doc, "The repository is https://github.com/onemoretime-hh/ethereum-cloud-blockchain-labs. The assignment code is contained in the assignment-02-casino-dapp directory, with separate contracts, tests, deployment scripts, frontend, and report folders.")

add_heading(doc, "Verification results", 2)
three_line_table(
    doc,
    ["Item", "Verified result", "Command or evidence"],
    [
        ("Solidity contract", "Compilation succeeded with Solidity 0.8.28.", "pnpm compile"),
        ("Contract tests", "10 automated tests passed.", "pnpm test"),
        ("Local deployment", "Deployment succeeded on the Hardhat simulated network.", "pnpm deploy:local"),
        ("React frontend", "TypeScript checking and Vite production build succeeded.", "pnpm frontend:build"),
    ],
    [1.35, 3.0, 2.45],
)
add_body(doc, "The local deployment address 0x5FbDB2315678afecb367f032d93F642f64180aa3 is evidence from the disposable Hardhat network only; it is not represented as a Sepolia address. A public Sepolia deployment requires a student-controlled test wallet and test ETH. The README gives the exact deployment command, after which the verified Sepolia address and transaction hash can be recorded. The built frontend/dist directory is ready to add to IPFS, and the resulting root CID can then be recorded for submission.")

add_heading(doc, "Design scope and security limitation", 2)
add_body(doc, "The implementation includes input validation, duplicate-bet prevention, pull-based prize withdrawal, and automated tests, but it remains a classroom example. In particular, the winning number hashes block.prevrandao, block.timestamp, the round number, and the contract address. A block producer or informed participant may influence or predict such inputs, so this method is not a secure source of randomness. A production system involving real value would require a verifiable randomness service, a fuller threat model, an audit, careful economic design, and compliance review.")

add_heading(doc, "Deployment procedure", 2)
add_body(doc, "First install dependencies, compile the contract, run the tests, and export the ABI. Next set SEPOLIA_RPC_URL and SEPOLIA_PRIVATE_KEY only in the local shell or Hardhat keystore and run the Sepolia deployment script. Copy the confirmed contract address into frontend/.env.local as VITE_CASINO_ADDRESS, build the React project, and test it in MetaMask on Sepolia. Finally add the complete frontend/dist directory to IPFS. If a custom domain is required, publish a DNSLink TXT record that points to the root CID. Private keys, seed phrases, RPC secrets, build caches, and local environment files must not be committed.")

doc.add_page_break()
add_heading(doc, "References", 1)
references = [
    "1. Course lecture, Ethereum DApps And Related Tools Primer, smart contract use cases slide.",
    "2. Ethereum.org, Ethereum Accounts, https://ethereum.org/developers/docs/accounts/.",
    "3. Ethereum.org, Transactions, https://ethereum.org/developers/docs/transactions/.",
    "4. Ethereum.org, Ethereum Virtual Machine, https://ethereum.org/developers/docs/evm/.",
    "5. Ethereum.org, Opcodes for the EVM, https://ethereum.org/developers/docs/evm/opcodes/.",
    "6. Circle, USDC Contract Addresses, https://developers.circle.com/stablecoins/usdc-contract-addresses.",
    "7. ENS Documentation, Deployments, https://docs.ens.domains/learn/deployments/.",
    "8. Aave Documentation, Governance, https://www.aave.com/docs/ecosystem/governance.",
    "9. Aave Documentation, Pool, https://aave.com/docs/aave-v3/smart-contracts/pool.",
    "10. Uniswap Documentation, Uniswap V2 Deployments, https://developers.uniswap.org/docs/protocols/v2/deployments.",
    "11. MetaMask Documentation, eth_accounts and eth_requestAccounts, https://docs.metamask.io/metamask-connect/evm/reference/json-rpc-api/eth_accounts/.",
    "12. Ethereum Foundation, Ropsten, Rinkeby and Kiln Deprecation Announcement, https://blog.ethereum.org/2022/06/21/testnet-deprecation.",
    "13. Hardhat Documentation, Configuration and Getting Started with Hardhat 3, https://hardhat.org/docs/reference/configuration and https://hardhat.org/docs/getting-started.",
]
for reference in references:
    add_body(doc, reference)

# Required standalone final paragraph.
final_paragraph = add_body(doc, "The original tutorial depended on the retired Ropsten network, the legacy MetaMask window.web3 injection API, and package versions whose interfaces have since changed. This submission therefore makes the necessary compatibility updates by using Sepolia, the current EIP-1193 MetaMask provider through ethers.js BrowserProvider, Solidity 0.8.x, Hardhat 3, and a current React and Vite toolchain while preserving the tutorial's classroom learning objectives.")
final_paragraph.paragraph_format.space_before = Pt(12)

for section in doc.sections:
    add_page_number(section.footer.paragraphs[0])

doc.core_properties.title = "Assignment Lab 2"
doc.core_properties.subject = "Ethereum smart contract exercises and Casino dApp"
doc.core_properties.author = "Jiahao Wan"
doc.core_properties.keywords = "Ethereum, Solidity, Casino dApp, Assignment 2"
doc.save(OUT)
print(OUT)
