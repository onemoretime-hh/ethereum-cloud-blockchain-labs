from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn


REPORT = Path(__file__).with_name("Assignment_02_Wan_Jiahao.docx")
assert REPORT.exists(), "Build the Word report first"
doc = Document(REPORT)
text = "\n".join(paragraph.text for paragraph in doc.paragraphs)

for label in ("1(a)", "1(b)", "1(c)", "2(a)", "2(b)", "2(c)", "3 Prevention", "4 Transactions", "5(a)", "5(b)", "5(c)"):
    assert label in text, f"missing answer label {label}"

for required in (
    "Wan, Jiahao",
    "2621751515@qq.com",
    "https://github.com/onemoretime-hh/ethereum-cloud-blockchain-labs",
    "retired Ropsten network",
    "legacy MetaMask window.web3",
):
    assert required in text or any(required in cell.text for table in doc.tables for row in table.rows for cell in row.cells)

addresses = [
    "0xA0b86991c6218b36c1d19D4a2e9Eb0cE3606eB48",
    "0x00000000000C2E074eC69A0dFb2997BA6C7d2e1e",
    "0x9AEE0B04504CeF83A65AC3f0e838D0593BCb2BC7",
    "0x7a250d5630B4cF539739dF2C5dAcb4c659F2488D",
    "0x87870Bca3F3fD6335C3F4ce8392D69350B4fA4E2",
]
table_text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
assert all(address in table_text for address in addresses)

final_nonempty = next(p for p in reversed(doc.paragraphs) if p.text.strip())
assert final_nonempty.text.startswith("The original tutorial depended on the retired Ropsten network")

for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        if run.text.strip() and run.font.color.rgb is not None:
            assert str(run.font.color.rgb) == "000000"

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    if run.text.strip() and run.font.color.rgb is not None:
                        assert str(run.font.color.rgb) == "000000"

with ZipFile(REPORT) as archive:
    xml = archive.read("word/document.xml").decode("utf-8")
assert "w:color w:val=\"000000\"" in xml
assert xml.count("w:tblBorders") >= 2
assert "w:insideV w:val=\"nil\"" in xml

print("Report structural checks passed")
